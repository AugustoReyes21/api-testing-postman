"""Genera el DOCX académico a partir de evidencias reales previamente revisadas."""
from pathlib import Path
import json
from datetime import datetime, timezone, timedelta
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
ART = ROOT / "artifacts"
SHOTS = ROOT / "evidence" / "screenshots"
ART.mkdir(exist_ok=True)

# Recortes de capturas originales; nunca se generan ni dibujan interfaces.
source = ROOT / ".runtime" / "postman-create-persist.png"
if source.exists():
    with Image.open(source) as im:
        im.crop((480, 620, 1540, 960)).save(SHOTS / "postman-create.png")
source = SHOTS / "postman-corrected.png"
if source.exists():
    with Image.open(source) as im:
        im.crop((0, 0, 1060, 300)).save(SHOTS / "postman-corrected-summary.png")
source = SHOTS / "postman-controlled-defect.png"
if source.exists():
    with Image.open(source) as im:
        im.crop((0, 0, 1060, 675)).save(SHOTS / "postman-controlled-defect-focus.png")

final = json.loads((ROOT / "evidence/runs/clean-repeat.json").read_text(encoding="utf-8"))
bad = json.loads((ROOT / "evidence/runs/controlled-defect.json").read_text(encoding="utf-8"))
verified = json.loads((ROOT / "evidence/verification.json").read_text(encoding="utf-8"))
assert final["success"] and verified["implementationRestored"] and verified["contractUnchanged"]
assert final["caseCount"] == 18 and final["assertionsFailed"] == 0
assert bad["assertionsFailed"] == 15

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
normal = doc.styles["Normal"]
normal.font.name, normal.font.size = "Arial", Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.space_after = Pt(7)
for name in ("Title", "Heading 1", "Heading 2"):
    style = doc.styles[name]
    style.font.name = "Arial"
    style.font.size = Pt(12 if name == "Title" else 11)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(8)
    style.paragraph_format.space_after = Pt(7)
doc.core_properties.author = ""
doc.core_properties.last_modified_by = ""
doc.core_properties.title = "Automatización de pruebas de API con Postman"
doc.core_properties.subject = "Práctica académica verificable"
doc.core_properties.comments = ""

def p(text):
    para = doc.add_paragraph(text)
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return para

def heading(text):
    return doc.add_paragraph(text, "Heading 1")

def image(name, width=6.5):
    path = SHOTS / name
    assert path.exists(), f"Falta evidencia: {path}"
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(5)
    para.add_run().add_picture(str(path), width=Inches(width))

def page():
    doc.add_page_break()

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = str(text)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = str(text)
    for row in t.rows:
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(3)
                para.paragraph_format.line_spacing = 1.0
                for run in para.runs:
                    run.font.name, run.font.size = "Arial", Pt(11)
    return t

doc.add_paragraph("Automatización de pruebas de API con Postman", "Title")
p("Estudiante: José Augusto Reyes Samayoa")
p("Carné: 2790-22-14075")
date = datetime.now(timezone(timedelta(hours=-6), name="America/Guatemala"))
months = "enero febrero marzo abril mayo junio julio agosto septiembre octubre noviembre diciembre".split()
p(f"Fecha de verificación: {date.day} de {months[date.month-1]} de {date.year}.")
heading("1. Introducción")
p("La práctica integra una API de pedidos y una colección que comprueba su comportamiento con datos ficticios. No basta con que el servidor responda: también se revisan los permisos, la persistencia, las reglas de entrada y el contrato de las respuestas. Las evidencias incluidas proceden de ejecuciones reales en Windows.")
heading("2. Objetivo")
p("Automatizar pruebas funcionales y de contrato con Postman, demostrar que detectan un defecto introducido de forma controlada y dejar un procedimiento que pueda repetirse desde una base de datos limpia.")
heading("3. API implementada")
p("Se utilizó Node.js 24.18.0, Express 5.3.0 y SQLite 3.53.1. Dos usuarios ficticios obtienen sesiones temporales y administran únicamente sus pedidos. Los importes se calculan en centavos enteros; el propietario se obtiene de la sesión. Los errores tienen código, mensaje, detalles y un identificador de correlación.")
table(["Operación", "Finalidad"], [
    ("POST /auth/login", "Autenticar al usuario ficticio."),
    ("POST /orders", "Crear un pedido protegido."),
    ("GET /orders y /orders/{id}", "Listar pedidos propios y consultar uno persistido."),
    ("PUT /orders/{id}/status", "Asignar un estado de forma idempotente."),
    ("DELETE /orders/{id}", "Eliminar un pedido y verificar su ausencia."),
])
p("La API escucha por defecto en http://127.0.0.1:3000. El comando npm run seed genera credenciales ficticias locales y restablece los datos. SQLite y el ambiente con contraseñas quedan excluidos de Git.")

page()
heading("4. Especificación y pruebas realizadas")
p("El contrato OpenAPI 3.1 documenta siete operaciones, autenticación, parámetros, encabezados, errores y reglas de validación. Swagger Parser y Redocly verificaron el documento. Ajv2020 comprobó los esquemas de los intercambios HTTP registrados. Las aserciones de cuerpos JSON de Postman se evaluaron por separado: no se presentan como una validación completa de OpenAPI.")
p("La matriz contiene 18 casos distintos: TC-01 a TC-12 cubren los escenarios obligatorios; TC-13 a TC-18 agregan contraseña incorrecta, paginación inválida, eliminación repetida, JSON malformado, asignación de propietario y aislamiento del listado. Los 46 pasos incluyen preparación y limpieza. Las 234 aserciones no se cuentan como 234 casos.")
p("El ambiente publicado se importó sin contraseñas ni tokens. Las variables necesarias se distinguen en la siguiente captura de Postman.")
image("postman-environment-example.png")
p("Para la ejecución se seleccionó el ambiente local con usuarios ficticios; sus contraseñas se muestran ocultas y no forman parte del repositorio.")
image("postman-environment-local.png")
p("Los IDs y tokens se capturan durante cada ejecución. La preparación crea los pedidos necesarios para la paginación y la limpieza elimina los recursos generados. La matriz detallada está en docs/test-matrix.md y docs/test-matrix.csv.")

page()
heading("5. Ejecución real en Postman")
p("La colección se importó y ejecutó con Collection Runner de Postman 12.32.4 instalado en el dispositivo. La captura muestra la API disponible, la autenticación válida y el resultado de la colección completa: 234 aserciones aprobadas, cero fallidas y cero errores.")
image("postman-final.png")
p("También se ejecutó Newman 6.2.3. La ejecución corregida y la repetición desde otro estado limpio aprobaron los mismos 18 casos. Los resultados reutilizables se conservan en evidence/runs. La exportación JSON del Runner gráfico estaba restringida por el plan de Postman; no se contrató ni activó un plan adicional.")

page()
heading("6. Creación, persistencia y validación")
p("Al crear el pedido se comprobó HTTP 201, el importe calculado, el propietario, el esquema y el encabezado Location. El identificador se obtuvo de la respuesta, sin reutilizar uno de otra ejecución.")
image("postman-create.png")
p("La consulta posterior recuperó el mismo recurso y aprobó la comparación completa con el pedido creado. Además se verificó la persistencia después de detener e iniciar la API con el mismo SQLite; ese resultado está en evidence/persistence-restart.json.")
image("postman-persistence.png")
p("La omisión de items produjo HTTP 400 y un error estructurado. Las pruebas verificaron su esquema y el detalle del campo obligatorio.")
image("postman-validation-error.png")

page()
heading("7. Permisos, paginación e idempotencia")
p("El usuario B recibió 403 al intentar consultar, modificar o borrar un pedido de A. Después, A recuperó su pedido sin cambios. La captura corresponde al intento de modificación rechazado.")
image("postman-access-control.png")
p("La paginación comprobó límites, páginas diferentes, ausencia de duplicados y metadatos. PUT cambió pending a confirmed una sola vez; la repetición y la consulta final conservaron versión, fecha e importe. El Runner aprobó la comprobación de que no hubo una segunda mutación.")
image("postman-pagination-idempotency-grid.png", width=4.7)
p("Los intercambios registrados por Newman confirmaron los siguientes valores de TC-11. La versión y la fecha se conservaron después de la primera aplicación.")
idem = [e for e in final["executions"] if e["name"].startswith("TC-11")]
table(["Operación", "Estado / versión", "updatedAt (UTC)"], [(label, e["body"]["status"] + " / " + str(e["body"]["version"]), e["body"]["updatedAt"]) for label, e in zip(("Primer PUT", "Segundo PUT", "GET posterior"), idem)])

page()
heading("8. Fallo controlado y corrección")
p("Después de establecer la referencia correcta, se cambió temporalmente totalCents de número a texto en la serialización de la API. OpenAPI siguió exigiendo un entero. Aunque las solicitudes devolvieron 200 o 201, el esquema detectó el tipo incorrecto; también falló la comparación del importe. Postman mostró 15 aserciones fallidas.")
image("postman-controlled-defect-focus.png")
p("Se restauró Number(row.total_cents) y se reinició el servidor. Sin modificar el contrato ni debilitar las aserciones, el Runner volvió a aprobar las 234 comprobaciones.")
image("postman-corrected-summary.png")
p("npm run verify reproduce el defecto, exige su detección y restaura el código. La comprobación SHA-256 confirmó que el contrato no cambió. El parche demostrativo se conserva inactivo.")

page()
heading("9. Resultados finales")
table(["Ejecución Newman", "Casos", "Aprobadas", "Fallidas"], [
    ("Referencia funcional", 18, 234, 0),
    ("Defecto controlado", 18, 219, 15),
    ("Corrección desde estado limpio", 18, 234, 0),
    ("Repetición desde otro estado limpio", 18, 234, 0),
])
p("Cada ejecución procesó 46 solicitudes. El validador adicional encontró 16 respuestas incompatibles durante el defecto y ninguna después de corregirlo. Postman de escritorio confirmó 234 aprobadas y cero fallidas en la ejecución final. Los informes diferencian resultados por ejecución; no suman repeticiones para aumentar el total.")
heading("10. Conclusiones")
p("Las pruebas detectaron tanto incumplimientos funcionales como una respuesta con tipo incorrecto, aun cuando su código HTTP era de éxito. Los intentos de acceso entre usuarios fueron rechazados y la consulta posterior confirmó que no modificaron el pedido. La idempotencia se comprobó comparando el recurso persistido, incluida su versión y su fecha.")
p("Las dos ejecuciones finales desde datos restablecidos terminaron sin fallos. La API también conservó un pedido después de reiniciar el proceso. Estos resultados respaldan la reproducibilidad de la práctica dentro del alcance de los escenarios automatizados.")
heading("11. Repositorio y reproducción")
p("https://github.com/AugustoReyes21/api-testing-postman")
p("El repositorio es público y el docente puede consultarlo sin solicitar permisos. Incluye código, contrato, matriz, colección, ambiente sin secretos, scripts, resultados, capturas y el DOCX editable. El PDF se entrega por separado y no se incluye en GitHub. Para verificar todo: npm ci y npm run verify. Para la aplicación gráfica: npm run seed, npm start e importar los archivos indicados en el README.")

# Sin secciones de encabezado/pie ni leyendas bajo imágenes.
output = ART / "informe-api-postman.docx"
doc.save(output)
print(f"DOCX generado: {output.name}")
